"""CSV/XLS/XLSX com cabeçalhos originais preservados."""

import csv
import io
from pathlib import Path

import xlrd
from openpyxl import load_workbook


def _xls_value(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> str:
    """Converte células BIFF sem transformar datas ou inteiros em valores ambíguos."""
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
        return ""
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate_as_datetime(cell.value, book.datemode)
        return value.isoformat(sep=" ")
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return str(int(cell.value))
    return str(cell.value).strip()


def _row(names: list[str], values: list[str]) -> dict[str, str]:
    """Combina cabeçalhos repetidos, preservando o último valor não vazio."""
    result: dict[str, str] = {}
    for name, value in zip(names, values):
        if name and (value or name not in result):
            result[name] = value
    return result


def read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = content.decode("cp1252")
        sample = decoded[:4096]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.DictReader(io.StringIO(decoded), delimiter=delimiter)
        if not reader.fieldnames or any(field is None or not field.strip() for field in reader.fieldnames):
            raise ValueError("Cabeçalhos vazios ou ausentes no CSV")
        return [{str(k).strip(): str(v or "").strip() for k, v in row.items()} for row in reader]
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        result: list[dict[str, str]] = []
        for sheet in workbook:
            iterator = sheet.iter_rows(values_only=True)
            try:
                header = next(iterator)
            except StopIteration:
                continue
            names = [str(cell or "").strip() for cell in header]
            if not any(names):
                continue
            for values in iterator:
                if any(value is not None and str(value).strip() for value in values):
                    result.append(_row(names, [str(value or "").strip() for value in values]))
        workbook.close()
        return result
    if suffix == ".xls":
        # As exportações BIFF do CAD repetem setores OLE. Excel as abre normalmente;
        # o modo tolerante do xlrd permite lê-las sem alterar o arquivo original.
        workbook = xlrd.open_workbook(
            file_contents=content,
            on_demand=True,
            ignore_workbook_corruption=True,
        )
        result: list[dict[str, str]] = []
        try:
            for sheet in workbook.sheets():
                if not sheet.nrows:
                    continue
                names = [_xls_value(workbook, sheet.cell(0, column))
                         for column in range(sheet.ncols)]
                if not any(names):
                    continue
                for row_index in range(1, sheet.nrows):
                    values = [_xls_value(workbook, sheet.cell(row_index, column))
                              for column in range(sheet.ncols)]
                    if any(values):
                        result.append(_row(names, values))
        finally:
            workbook.release_resources()
        return result
    raise ValueError("Formato aceito: CSV, XLS ou XLSX")
